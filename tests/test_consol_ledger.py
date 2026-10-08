"""Offline ledger and report email integration: no database or real email calls."""
import base64
import datetime
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook
from api import email_service, ledger_service, main, pdf_service
import weekly_report_mailer as mailer


class ConsolLedgerTest(unittest.TestCase):
    def test_standard_sea_ledger_contains_query_rows_before_dashboard_normalization(self):
        records = [{"Console_Number": "SEA-1", "Master_Airway_Bill": "00001", "FCL_TEU_Count": 2.5,
                    "LCL_Volume": 10, "Revenue_USD": 100, "Source_Column": "first"},
                   {"Console_Number": "SEA-1", "Master_Airway_Bill": "00001", "FCL_TEU_Count": 2.5,
                    "LCL_Volume": 10, "Revenue_USD": 200, "Source_Column": "second"}]
        request = MagicMock(headers={"X-Consol-Ledger": "true"})
        with patch.object(main, "get_sea_source_data", return_value=records) as query:
            result = main.fetch_data("2026-09-21", "2026-09-27", transport_mode="SEA", request=request)
        query.assert_called_once()
        self.assertEqual(result["ledger_records"], records)
        self.assertEqual(len(result["data"]), 1)
        self.assertNotIn("Source_Column", result["data"][0])
        self.assertEqual(result["ledger_records"][1]["Source_Column"], "second")

    def test_custom_sea_ledger_keeps_original_query_columns(self):
        records = [{"Console_Number": "SEA-1", "Master_Airway_Bill": "00001", "FCL_TEU_Count": 2.5,
                    "LCL_Volume": 10, "Revenue_USD": 100, "Source_Column": "source value"}]
        with patch.object(main, "execute_custom_query", return_value=records) as query:
            result = main.custom_query(main.CustomQueryRequest(query="SELECT * WHERE TransportMode = 'SEA'",
                transport_mode="SEA"), request=MagicMock(headers={"X-Consol-Ledger": "true"}))
        query.assert_called_once()
        self.assertEqual(result["ledger_records"], records)
        self.assertNotIn("Source_Column", result["data"][0])
        self.assertNotIn("Airline", result["ledger_records"][0])

    def test_export_keeps_every_row_column_type_and_duplicate_source_record(self):
        rows = [{"Console_Number": "000123", "FCL_TEU_Count": 2.5, "LCL_Volume": 10.25,
                 "Revenue_USD": 1000, "ETD": "2026-09-21T00:15:00+05:30", "Note": "=1+1"}
                for _ in range(125)]
        rows[-1]["Extra_Source_Column"] = "last row retained"
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sea.xlsx"
            ledger_service.generate_consol_ledger(path, rows, "SEA")
            book = load_workbook(path)
            sheet = book["Consol Ledger"]
            self.assertEqual(sheet.max_row, 129)
            self.assertEqual([cell.value for cell in sheet[4]], list(rows[-1]))
            self.assertEqual(sheet["A5"].value, "000123")
            self.assertEqual(sheet["A5"].data_type, "s")
            self.assertEqual(sheet["B5"].value, 2.5)
            self.assertEqual(sheet["B5"].data_type, "n")
            self.assertEqual(sheet["E5"].value, rows[0]["ETD"])
            self.assertEqual(sheet["F5"].value, "=1+1")
            self.assertEqual(sheet["F5"].data_type, "s")
            self.assertEqual(sheet["G129"].value, "last row retained")
            self.assertEqual(sheet.freeze_panes, "B5")
            self.assertEqual(sheet.tables["ConsolRecords"].ref, "A4:G129")
            book.close()

    def test_empty_result_still_exports_valid_ledger(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "empty.xlsx"
            ledger_service.generate_consol_ledger(path, [], "AIR")
            book = load_workbook(path)
            self.assertIn("No records", book.active["A4"].value)
            book.close()

    def test_air_source_dates_and_shipment_columns_are_retained(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "air.xlsx"
            row = {"Console_Number": "AIR-1", "ShipmentNumber": "00001", "Total_Tonnage": 2500,
                   "ETD": datetime.datetime(2026, 9, 21, 10, 30), "Linked": True, "Missing": None}
            ledger_service.generate_consol_ledger(path, [row])
            book = load_workbook(path)
            self.assertEqual([cell.value for cell in book.active[5]], list(row.values()))
            book.close()

    def mock_graph(self):
        app = MagicMock()
        app.acquire_token_for_client.return_value = {"access_token": "fixture-token"}
        return app

    def test_manual_and_scheduled_air_sea_both_sends_have_matching_excel_and_cleanup(self):
        for scheduled in (False, True):
            for mode in ("AIR", "SEA", "BOTH"):
                with self.subTest(scheduled=scheduled, mode=mode):
                    created = []
                    def generate(output_path, transport_mode, **kwargs):
                        created.append(output_path)
                        Path(output_path).write_bytes(b"fixture PDF")
                        ledger_service.generate_consol_ledger(ledger_service.ledger_path(output_path),
                            [{"Console_Number": transport_mode + "-1", "Source_Value": 123}], transport_mode)
                    config = {"is_active": 1, "recipient_email": "fixture@example.test", "frequency": "weekly",
                              "filters": {"transport_mode": mode, "country": "India", "company_code": "IND"}}
                    with patch.dict("os.environ", {"DB_SERVER": "fixture", "DB_USER": "fixture", "DB_PASSWORD": "fixture",
                            "MAIL_AZURE_TENANT_ID": "fixture", "MAIL_AZURE_CLIENT_ID": "fixture",
                            "MAIL_AZURE_CLIENT_SECRET": "fixture", "SENDER_EMAIL": "fixture@example.test", "VERCEL": ""}), \
                         patch.object(main, "generate_dashboard_pdf", side_effect=generate), \
                         patch.object(main, "get_schedule", return_value=config), \
                         patch.object(main, "get_supabase_stations", return_value=[]), \
                         patch.object(main, "get_supabase_branches", return_value=[]), \
                         patch.object(email_service, "ConfidentialClientApplication", return_value=self.mock_graph()), \
                         patch.object(email_service.requests, "post", return_value=MagicMock(status_code=202)) as post, \
                         patch.object(email_service, "log_email_transaction"):
                        if scheduled:
                            main.execute_scheduled_report_job("fixture")
                        else:
                            main.send_report(main.ReportRequest(start_date="2026-09-21", end_date="2026-09-27",
                                recipient_email="fixture@example.test", country="India", transport_mode=mode))
                        post.assert_called_once()
                        attachments = post.call_args.kwargs["json"]["message"]["attachments"]
                        self.assertEqual(len(attachments), 4 if mode == "BOTH" else 2)
                        for pdf, excel in zip(attachments[::2], attachments[1::2]):
                            self.assertEqual(pdf["contentType"], "application/pdf")
                            self.assertEqual(excel["name"], ledger_service.ledger_attachment_name(pdf["name"]))
                            self.assertEqual(excel["contentType"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                            book = load_workbook(io.BytesIO(base64.b64decode(excel["contentBytes"])))
                            self.assertEqual(book.active["A5"].value, "SEA-1" if pdf["name"].startswith("Sea_") else "AIR-1")
                            book.close()
                    for path in created:
                        self.assertFalse(Path(path).exists())
                        self.assertFalse(Path(ledger_service.ledger_path(path)).exists())

    def test_missing_excel_prevents_graph_send(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.pdf"
            path.write_bytes(b"fixture PDF")
            with patch.dict("os.environ", {"MAIL_AZURE_TENANT_ID": "fixture", "MAIL_AZURE_CLIENT_ID": "fixture",
                    "MAIL_AZURE_CLIENT_SECRET": "fixture", "SENDER_EMAIL": "fixture@example.test"}), \
                 patch.object(email_service, "ConfidentialClientApplication", return_value=self.mock_graph()), \
                 patch.object(email_service.requests, "post") as post, \
                 patch.object(email_service, "log_email_transaction"):
                with self.assertRaises(FileNotFoundError):
                    email_service.send_pdf_via_graph(pdf_path=str(path), recipient_email="fixture@example.test")
                post.assert_not_called()

    def test_standalone_mailer_delegates_to_bundle_sender(self):
        with patch.object(email_service, "send_pdf_via_graph") as send:
            mailer.send_email_via_graph("fixture.pdf", "India", "2026-09-21", "2026-09-27",
                                       ["fixture@example.test"], "SEA")
        self.assertEqual(send.call_args.kwargs["pdf_path"], "fixture.pdf")
        self.assertIn("Consol Ledger Excel", send.call_args.kwargs["body"])

    def test_bundle_requests_ledger_from_same_pdf_generation(self):
        with patch.object(pdf_service, "generate_dashboard_pdf") as generate:
            pdf_service.generate_report_bundle("fixture.pdf", transport_mode="SEA", max_data_rows=1)
        self.assertEqual(generate.call_args.kwargs["ledger_output_path"], "fixture.xlsx")
        self.assertEqual(generate.call_args.kwargs["max_data_rows"], 1)


if __name__ == "__main__":
    unittest.main()
