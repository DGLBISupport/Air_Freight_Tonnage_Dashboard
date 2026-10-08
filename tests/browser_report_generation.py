"""Render real AIR/SEA PDFs with fixture APIs; never send an email."""
import json
import time
from pathlib import Path
from urllib.parse import urlencode, urlparse, parse_qs
from unittest.mock import patch

from playwright.sync_api import sync_playwright
from api import pdf_service
from openpyxl import load_workbook
from pypdf import PdfReader
from browser_freight_modes import BASE, intercept, respond, requests_seen, air_rows, sea_rows


def main():
    Path("outputs").mkdir(exist_ok=True)
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(channel="chrome", headless=True)
        api_hosts = []
        def fixtures(route):
            parsed = urlparse(route.request.url)
            if parsed.path.startswith("/api/"):
                api_hosts.append(parsed.hostname)
            if (route.request.headers.get("x-consol-ledger") == "true"
                    and parsed.path in ("/api/data", "/api/custom-query")):
                body = route.request.post_data_json if route.request.method == "POST" else None
                params = parse_qs(parsed.query)
                mode = body.get("transport_mode") if body else params.get("transport_mode", ["AIR"])[0]
                records = sea_rows if mode == "SEA" else air_rows
                requests_seen.append((parsed.path, params, body))
                respond(route, {"status": "success", "data": records,
                    "ledger_records": [dict(row, Original_Query_Column="source record") for row in records]})
            elif parsed.path.endswith("/expired"):
                route.fulfill(status=404, content_type="application/json",
                              body=json.dumps({"detail": "Query not found or expired"}))
            else:
                intercept(route)

        # Verify a production build viewed locally uses local API URLs.
        page = browser.new_page()
        page.route("**/*", fixtures)
        started = time.monotonic()
        page.goto(BASE + "/print-view/?" + urlencode({"transport_mode": "SEA", "include_sector_distribution": "false"}))
        page.locator("#pdf-ready").wait_for()
        elapsed = time.monotonic() - started
        assert api_hosts and all(host in ("localhost", "127.0.0.1") for host in api_hosts), api_hosts
        assert not any(path == "/api/sector-carrier-distribution" for path, _, _ in requests_seen)
        print(f"PASS: local SEA preview ready in {elapsed:.2f}s with fixture data; excluded sectors were not requested.")

        page.goto(BASE + "/print-view/?mode=custom-sql&query_id=expired&transport_mode=SEA")
        page.locator("#print-error").wait_for()
        assert page.locator("#pdf-ready").count() == 0
        assert "expired" in page.locator("#print-error").inner_text()
        page.close()

        # Exercise the application's real PDF function using installed Chrome
        # and intercept every API/external request from its fresh browser pages.
        class FixtureBrowser:
            page = None
            def new_page(self, **kwargs):
                self.page = browser.new_page(**kwargs)
                self.page.route("**/*", fixtures)
                return self.page
            def close(self):
                if self.page:
                    self.page.close()

        class FixtureChromium:
            def launch(self, **kwargs):
                return FixtureBrowser()

        class FixturePlaywright:
            chromium = FixtureChromium()
            def __enter__(self):
                return self
            def __exit__(self, *args):
                return False

        for mode in ("AIR", "SEA"):
            start = len(requests_seen)
            started = time.monotonic()
            output = f"outputs/{mode.lower()}-checked-report.pdf"
            with patch.object(pdf_service, "sync_playwright", return_value=FixturePlaywright()), \
                 patch.object(pdf_service, "get_tonnage_base_url", return_value=BASE):
                pdf_service.generate_report_bundle(output, start_date="2026-09-21", end_date="2026-09-27",
                    transport_mode=mode, mode="custom-sql", query_id="expired",
                    custom_sql=f"SELECT * FROM vt WHERE TransportMode = '{mode}'", include_sector_distribution=False)
            assert Path(output).stat().st_size > 20000
            assert all(page.extract_text().strip() for page in PdfReader(output).pages), "Report has a blank PDF page"
            workbook = load_workbook(Path(output).with_suffix(".xlsx"))
            sheet = workbook["Consol Ledger"]
            source_request = [item for item in requests_seen[start:] if item[0] == "/api/custom-query"]
            assert len(source_request) == 1, "Excel must reuse the PDF response without another query"
            expected = [dict(row, Original_Query_Column="source record") for row in (sea_rows if mode == "SEA" else air_rows)]
            assert sheet.max_row == len(expected) + 4
            columns = [cell.value for cell in sheet[4]]
            exported = [dict(zip(columns, row)) for row in sheet.iter_rows(min_row=5, values_only=True)]
            assert exported == expected, "Excel must retain all fetched values, fields and rows"
            assert sheet.freeze_panes == "B5"
            workbook.close()
            sent = requests_seen[start:]
            assert not any(path.startswith("/api/get-cached-query") for path, _, _ in sent)
            assert any(path == "/api/custom-query" and body["transport_mode"] == mode for path, _, body in sent)
            print(f"PASS: actual {mode} PDF generated in {time.monotonic() - started:.2f}s with fixture data; injected SQL bypassed the expired query ID.")
        # Standard reports must also retain more than the PDF's display limit.
        for mode, source in (("AIR", air_rows), ("SEA", sea_rows)):
            original = list(source)
            try:
                source[:] = [dict(original[i % len(original)], Console_Number=f"{mode}-{i:04}") for i in range(125)]
                start = len(requests_seen)
                output = f"outputs/{mode.lower()}-standard-checked-report.pdf"
                with patch.object(pdf_service, "sync_playwright", return_value=FixturePlaywright()), \
                     patch.object(pdf_service, "get_tonnage_base_url", return_value=BASE):
                    pdf_service.generate_report_bundle(output, start_date="2026-09-21", end_date="2026-09-27",
                        country="India", company_code="IND", transport_mode=mode, max_data_rows=1,
                        include_sector_distribution=False)
                book = load_workbook(Path(output).with_suffix(".xlsx"))
                sheet = book["Consol Ledger"]
                assert sheet.max_row == 129
                columns = [cell.value for cell in sheet[4]]
                exported = [dict(zip(columns, row)) for row in sheet.iter_rows(min_row=5, values_only=True)]
                assert exported == [dict(row, Original_Query_Column="source record") for row in source]
                assert len([item for item in requests_seen[start:] if item[0] == "/api/data"]) == 1
                book.close()
                print(f"PASS: standard {mode} Excel retains all 125 fetched records with PDF max_data_rows=1, without another query.")
            finally:
                source[:] = original
        browser.close()
        print("PASS: expired previews show an error and never become PDF-ready. No emails sent.")


if __name__ == "__main__":
    main()
