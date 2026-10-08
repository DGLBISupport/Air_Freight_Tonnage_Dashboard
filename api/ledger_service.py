"""Export the complete report query result, without PDF grouping or row limits.

The deployed Python API uses openpyxl; the Codex-only artifact runtime is not
available in the application container or users' Python virtual environments.
"""
import datetime
import json
import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def ledger_path(pdf_path):
    return str(Path(pdf_path).with_suffix(".xlsx"))


def ledger_attachment_name(pdf_name):
    return f"{Path(pdf_name).stem}_Consol_Ledger.xlsx"


def cleanup_report_files(*pdf_paths):
    """Remove both parts of temporary report bundles, including partial failures."""
    for pdf_path in pdf_paths:
        if pdf_path:
            for path in (pdf_path, ledger_path(pdf_path)):
                try:
                    Path(path).unlink(missing_ok=True)
                except OSError:
                    logging.getLogger(__name__).warning("Could not remove temporary report file %s", path)


def generate_consol_ledger(output_path, records, transport_mode="AIR"):
    if not isinstance(records, list) or any(not isinstance(row, dict) for row in records):
        raise ValueError("Consol Ledger requires the fetched report records.")
    columns = list(dict.fromkeys(key for row in records for key in row))
    if len(records) > 1048572 or len(columns) > 16384:
        raise ValueError("The full Consol Ledger exceeds Excel worksheet limits.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Consol Ledger"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = "3182CE" if transport_mode == "AIR" else "0D9488"
    sheet["A1"] = f"{transport_mode.title()} Freight Consol Ledger"
    sheet["A1"].font = Font(name="Arial", size=14, bold=True, color="1E293B")
    sheet["A2"] = f"Report query result: {len(records):,} records. All fetched rows and columns; no PDF row limit."
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color="64748B")
    if not columns:
        sheet["A4"] = "No records returned for this report."
        sheet.column_dimensions["A"].width = 85
    else:
        widths = [max(18, len(str(column)) + 3) for column in columns]
        for index, column in enumerate(columns, 1):
            cell = sheet.cell(4, index, str(column))
            cell.data_type = "s"
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="334155")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[4].height = 32
        for row_index, record in enumerate(records, 5):
            for col_index, column in enumerate(columns, 1):
                value = record.get(column)
                if isinstance(value, (dict, list)):
                    value = json.dumps(value, ensure_ascii=False)
                # Preserve timezone-bearing dates as source text; Excel cannot
                # store timezone-aware dates without discarding their offset.
                if isinstance(value, datetime.datetime) and value.tzinfo:
                    value = value.isoformat()
                if isinstance(value, str) and len(value) > 32767:
                    raise ValueError(f"{column} contains text exceeding Excel's cell limit.")
                cell = sheet.cell(row_index, col_index, value)
                cell.font = Font(name="Arial", size=10, color="1E293B")
                cell.alignment = Alignment(vertical="center")
                if isinstance(value, str):
                    # Query text is data, including identifiers beginning with =.
                    cell.data_type = "s"
                    cell.number_format = "@"
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif isinstance(value, (datetime.datetime, datetime.date)):
                    cell.number_format = "yyyy-mm-dd hh:mm:ss" if isinstance(value, datetime.datetime) else "yyyy-mm-dd"
                elif isinstance(value, (int, float)) and not isinstance(value, bool):
                    cell.number_format = ('#,##0.00' if "USD" in str(column).upper() else
                                          '#,##0' if isinstance(value, int) else '#,##0.########')
                widths[col_index - 1] = max(widths[col_index - 1], len(str(value)) + 2 if value is not None else 0)
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(width, 65)
        last_column = get_column_letter(len(columns))
        table = Table(displayName="ConsolRecords", ref=f"A4:{last_column}{len(records) + 4}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
        sheet.freeze_panes = "B5"
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(output_path)
    finally:
        workbook.close()
